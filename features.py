"""Landed cost, the import document checklist, and extra Excel sheets."""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from models import to_number

BASES = ["Value", "Quantity", "CBM or weight"]


def build_landed_cost(items: pd.DataFrame, freight: float, local_costs: float,
                      basis: str = "Value") -> tuple[pd.DataFrame, list[str]]:
    """items needs the columns: Description, Quantity, Value, Duty %, CBM or weight.
    Freight + insurance is added to the value (CIF), duty is charged on the CIF value,
    local costs (clearance, transport ...) are added after duty. Returns (table, notes)."""
    df = items.copy()
    for c in ["Quantity", "Value", "Duty %", "CBM or weight"]:
        df[c] = [to_number(v) for v in df[c]]
    notes: list[str] = []

    key = {"Value": "Value", "Quantity": "Quantity", "CBM or weight": "CBM or weight"}[basis]
    weights = df[key].fillna(0.0)
    total_w = float(weights.sum())
    if (freight or local_costs) and total_w <= 0:
        raise ValueError(f"Cannot share the costs by '{basis}': fill in that column for the items.")
    if (freight or local_costs) and basis == "CBM or weight" and df[key].isna().any():
        notes.append("Some items have no CBM or weight, so they get no share of freight or local costs.")
    share = (weights / total_w) if total_w > 0 else weights * 0.0

    value = df["Value"].fillna(0.0)
    out = pd.DataFrame({
        "Description": df["Description"],
        "Quantity": df["Quantity"],
        "Value": value,
        "Freight + insurance share": (share * freight).round(2),
    })
    out["CIF value"] = (out["Value"] + out["Freight + insurance share"]).round(2)
    if df["Duty %"].isna().any():
        notes.append(f"Duty % is missing for {int(df['Duty %'].isna().sum())} item(s): 0% was assumed. "
                     "Edit the Duty % column if you know the rate.")
    duty_pct = df["Duty %"].fillna(0.0)
    out["Duty %"] = duty_pct
    out["Duty amount"] = (out["CIF value"] * duty_pct / 100).round(2)
    out["Local costs share"] = (share * local_costs).round(2)
    out["Landed cost"] = (out["CIF value"] + out["Duty amount"] + out["Local costs share"]).round(2)
    qty = df["Quantity"]
    out["Landed cost per unit"] = [round(l / q, 2) if (q and q > 0) else None for l, q in zip(out["Landed cost"], qty)]
    out["Uplift %"] = [round((l - v) / v * 100, 1) if v else None for l, v in zip(out["Landed cost"], out["Value"])]
    if qty.isna().any():
        notes.append("Some items have no quantity, so their cost per unit is blank.")
    return out, notes


def checklist_table(requirements: dict[str, dict]) -> pd.DataFrame:
    """One row per document: HS code, agency, website, document. Codes with no agency say so."""
    rows = []
    for code, d in requirements.items():
        if not d.get("found"):
            rows.append({"HS code": code, "Description": "", "Duty %": None, "Agency": "",
                         "Document": f"Not available: {d.get('error') or 'no detail'}", "Website": ""})
            continue
        base = {"HS code": code, "Description": d.get("final_description", ""),
                "Duty %": to_number(d.get("rate"))}
        agencies = d.get("agencies") or []
        if not agencies:
            rows.append({**base, "Agency": "None (standard clearance)", "Document": "", "Website": ""})
        for a in agencies:
            docs = a.get("documents") or [""]
            for doc in docs:
                rows.append({**base, "Agency": a.get("name_en") or a.get("name_ar") or "", "Document": doc,
                             "Website": a.get("website") or ""})
    return pd.DataFrame(rows, columns=["HS code", "Description", "Duty %", "Agency", "Document", "Website"])


def append_sheets(xlsx_bytes: bytes, sheets: dict[str, pd.DataFrame]) -> bytes:
    """Add extra sheets (name -> table) to an Excel file that is already built."""
    wb = load_workbook(BytesIO(xlsx_bytes))
    for title, df in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append([str(c) for c in df.columns])
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDE6F2")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()