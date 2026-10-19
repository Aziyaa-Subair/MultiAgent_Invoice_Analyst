import re
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

from models import InvoiceData, find_col, to_number

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# optional invoice-level columns that can be added in front of the item rows
HEADER_OPTIONS = {
    "Invoice number": lambda i: i.invoice_number,
    "Invoice date": lambda i: i.invoice_date.isoformat() if i.invoice_date else "",
    "Supplier": lambda i: i.supplier_name,
    "Customer": lambda i: i.customer_name or "",
    "Currency": lambda i: i.currency or "",
    "Invoice total": lambda i: i.total,
}

HS_COLS = {  # key in the lookup result -> column name shown in the table / Excel
    "hs_matched_code": "HS lookup: matched code",
    "hs_description": "HS lookup: description",
    "hs_duty_pct": "HS lookup: duty %",
    "hs_uom": "HS lookup: unit",
    "hs_approval_required": "HS lookup: approval required",
    "hs_approval_agencies": "HS lookup: approval agencies",
    "hs_note": "HS lookup: note",
}
_HS_ALIASES = {"hs", "hs code", "hscode", "h s code", "hs no", "hs number", "hts", "hts code", "hsn", "hsn code",
               "tariff", "tariff code", "tariff no", "customs code", "commodity code"}
_CODE_WORDS = {"code", "hs", "hsn", "hts", "no", "number", "id", "sku", "ref", "tariff", "barcode", "sl", "sr"}


# ---------------------------------------------------------------- one invoice -> one table, headers as printed
def _words(text: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", str(text).lower()))


def unique_headers(columns: list[str]) -> list[str]:
    """Same names as printed; only a repeated or empty header gets a number so a table can hold it."""
    seen: dict[str, int] = {}
    out = []
    for n, c in enumerate(columns, 1):
        name = (c or "").strip() or f"Column {n}"
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return out


def table_frame(inv: InvoiceData) -> pd.DataFrame:
    """The invoice's own item table: its own column names, all values as printed (text)."""
    cols = unique_headers(inv.table.columns)
    n = len(cols)
    rows = [([str(c).strip() for c in r.cells] + [""] * n)[:n] for r in inv.table.rows]
    return pd.DataFrame(rows, columns=cols, dtype="object")


def role_column(inv: InvoiceData, header: str | None) -> str | None:
    i = find_col(inv.table.columns, header)
    names = unique_headers(inv.table.columns)
    return names[i] if i is not None and i < len(names) else None


def hs_column(inv: InvoiceData, frame: pd.DataFrame) -> str | None:
    col = role_column(inv, inv.table.hs_code_column)
    if col in frame.columns:
        return col
    for c in frame.columns:
        if _words(c) in _HS_ALIASES:
            return c
    return None


# ---------------------------------------------------------------- checks (on the possibly edited tables)
def find_issues(frames: dict[str, pd.DataFrame], results: list[tuple[str, InvoiceData]]) -> list[str]:
    issues = []
    for name, inv in results:
        df = frames[name]
        amount = role_column(inv, inv.table.amount_column)
        items_sum = None
        if amount in df.columns:
            nums = [n for n in (to_number(v) for v in df[amount]) if n is not None]
            items_sum = round(sum(nums), 2) if nums else None
        if not inv.totals_match(items_sum=items_sum):
            issues.append(f"{name}: totals do not add up (item amounts {items_sum}, subtotal {inv.subtotal}, "
                          f"tax {inv.tax}, total {inv.total}).")
        q = role_column(inv, inv.table.quantity_column)
        p = role_column(inv, inv.table.unit_price_column)
        if q in df.columns and p in df.columns and amount in df.columns:
            for pos, (qv, pv, av) in enumerate(zip(df[q], df[p], df[amount]), 1):
                qn, pn, an = to_number(qv), to_number(pv), to_number(av)
                if None not in (qn, pn, an) and abs(qn * pn - an) > 0.05:
                    issues.append(f"{name} row {pos}: {q} x {p} is not equal to {amount} "
                                  f"(check for a discount, VAT in the amount, or a misread number).")
    by_number: dict[str, list[str]] = {}
    for name, inv in results:
        by_number.setdefault(inv.invoice_number, []).append(name)
    for number, files in by_number.items():
        if len(files) > 1:
            issues.append(f"Invoice number {number} appears in {len(files)} files: {', '.join(files)}.")
    return issues


# ---------------------------------------------------------------- HS code enrichment
def hs_digits(value) -> str:
    return "" if value is None or (isinstance(value, float) and pd.isna(value)) else re.sub(r"\D", "", str(value))


def hs_codes_in(frames: dict[str, pd.DataFrame], results: list[tuple[str, InvoiceData]]) -> list[str]:
    codes: list[str] = []
    for name, inv in results:
        col = hs_column(inv, frames[name])
        if col is None:
            continue
        for v in frames[name][col]:
            d = hs_digits(v)
            if d and d not in codes:
                codes.append(d)
    return codes


def add_hs_columns(frame: pd.DataFrame, hs_col: str | None, hs_info: dict[str, dict]) -> pd.DataFrame:
    """Append the portal's answer right after the invoice's HS code column."""
    if not hs_info or hs_col is None or hs_col not in frame.columns:
        return frame
    out = frame.copy()
    infos = [hs_info.get(hs_digits(v)) for v in out[hs_col]]
    new_cols = list(HS_COLS.values())
    for key, col in HS_COLS.items():
        out[col] = [(i or {}).get(key) for i in infos]
    base = [c for c in out.columns if c not in new_cols]
    at = base.index(hs_col) + 1
    return out[base[:at] + new_cols + base[at:]]


# ---------------------------------------------------------------- Excel
def _protected(inv: InvoiceData, frame: pd.DataFrame) -> set[str]:
    p = set()
    col = hs_column(inv, frame)
    if col:
        p.add(col)
    return p


def convert_numeric(df: pd.DataFrame, protect: set[str]) -> pd.DataFrame:
    """Text that is a number in every filled cell becomes a real number (codes and serial numbers stay text)."""
    out = df.copy()
    for c in out.columns:
        if c in protect or set(_words(c).split()) & _CODE_WORDS:
            continue
        filled = [v for v in out[c] if v is not None and not (isinstance(v, float) and pd.isna(v)) and str(v).strip() != ""]
        if filled and all(to_number(v) is not None for v in filled):
            out[c] = [to_number(v) if (v is not None and str(v).strip() != "") else None for v in out[c]]
    return out


def details_table(results: list[tuple[str, InvoiceData]]) -> pd.DataFrame:
    rows = []
    for file, inv in results:
        row = {
            "Source file": file, "Invoice number": inv.invoice_number,
            "Date": inv.invoice_date.isoformat() if inv.invoice_date else "",
            "Supplier": inv.supplier_name, "Customer": inv.customer_name or "",
            "Currency": inv.currency or "", "Subtotal": inv.subtotal, "Tax": inv.tax,
            "Total": inv.total, "Items": len(inv.table.rows),
            "Totals match": "Yes" if inv.totals_match() else "No",
        }
        for ev in inv.other_fields:
            key = (ev.name or "").strip()
            if key and key not in row:
                row[key] = ev.value
        rows.append(row)
    return pd.DataFrame(rows)


def _style(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE6F2")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", Path(name).stem)[:28] or "Sheet"
    title, n = base, 2
    while title in used or title == "Invoice details":
        title, n = f"{base[:26]}_{n}", n + 1
    used.add(title)
    return title


def _with_info(sub: pd.DataFrame, inv: InvoiceData, file: str, header_info: list[str], source: bool) -> pd.DataFrame:
    sub = sub.copy()
    pos = 0
    if source:
        sub.insert(pos, "Source file", file)
        pos += 1
    for label in header_info:
        if label in HEADER_OPTIONS:
            sub.insert(pos, label, HEADER_OPTIONS[label](inv))
            pos += 1
    return sub


def build_excel(frames: dict[str, pd.DataFrame], selected: dict[str, list[str]],
                results: list[tuple[str, InvoiceData]], layout: str = "one_sheet",
                header_info: list[str] | None = None) -> bytes:
    """frames = each file's table (its own column names); selected = the columns chosen per file.
    One file -> sheets 'Items' + 'Invoice details'.
    Several files -> layout 'one_sheet': 'All items' (a 'Source file' column holds the PDF name; columns an
    invoice does not have stay blank) or 'per_pdf': one sheet per PDF with only that invoice's own columns."""
    header_info = header_info or []
    prepared: dict[str, pd.DataFrame] = {}
    for file, inv in results:
        df = frames[file]
        cols = [c for c in df.columns if c in selected.get(file, [])]
        prepared[file] = convert_numeric(df[cols], _protected(inv, df))
    details = details_table(results)

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        if len(results) == 1:
            file, inv = results[0]
            _with_info(prepared[file], inv, file, header_info, source=False).to_excel(w, sheet_name="Items", index=False)
            pd.DataFrame({
                "field": list(details.columns),
                "value": ["" if pd.isna(details.iloc[0][c]) else details.iloc[0][c] for c in details.columns],
            }).to_excel(w, sheet_name="Invoice details", index=False)
        else:
            if layout == "per_pdf":
                used: set[str] = set()
                for file, inv in results:
                    prepared[file].to_excel(w, sheet_name=_sheet_title(file, used), index=False)
            else:
                parts, display = [], {}
                for file, inv in results:
                    sub = _with_info(prepared[file], inv, file, header_info, source=True)
                    renamed = []
                    for c in sub.columns:
                        renamed.append(display.setdefault(_words(c) or c, c))
                    sub.columns = renamed
                    parts.append(sub)
                pd.concat(parts, ignore_index=True, sort=False).to_excel(w, sheet_name="All items", index=False)
            details.to_excel(w, sheet_name="Invoice details", index=False)
        for ws in w.sheets.values():
            _style(ws)
    return buf.getvalue()


# ---------------------------------------------------------------- helpers used by main.py / pipeline.py / evaluate.py
def safe_stem(file_name: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(file_name).stem).strip("_")
    return stem or "invoice"


def excel_name(file_name: str) -> str:
    return f"{safe_stem(file_name)}.xlsx"


def invoice_xlsx_bytes(inv: InvoiceData, file_name: str = "invoice.pdf") -> bytes:
    """One workbook for ONE invoice with all of its own columns."""
    frame = table_frame(inv)
    return build_excel({file_name: frame}, {file_name: list(frame.columns)}, [(file_name, inv)])


def export_to_excel(invoices: list[InvoiceData], out_path: str) -> None:
    """Combined workbook used by pipeline.py and evaluate.py (sheets: Invoices, Line Items)."""
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    headers = pd.DataFrame([
        {
            "invoice_number": i.invoice_number, "date": i.invoice_date,
            "supplier": i.supplier_name, "customer": i.customer_name,
            "currency": i.currency, "subtotal": i.subtotal, "tax": i.tax,
            "total": i.total, "totals_match": i.totals_match(),
        }
        for i in invoices
    ])
    parts = []
    for i in invoices:
        f = table_frame(i)
        f.insert(0, "invoice_number", i.invoice_number)
        parts.append(f)
    lines = pd.concat(parts, ignore_index=True, sort=False) if parts else pd.DataFrame({"invoice_number": []})
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        headers.to_excel(writer, sheet_name="Invoices", index=False)
        lines.to_excel(writer, sheet_name="Line Items", index=False)